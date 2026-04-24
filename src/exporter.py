"""
exporter.py — Convert LLM results to JSON strings or styled Excel files.
"""

import json
import pathlib
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def to_json(result_text: str) -> str:
    """Return a pretty-printed JSON string.

    If result_text is already valid JSON, it is re-formatted.
    Otherwise it is wrapped in {"result": ...}.
    """
    try:
        parsed = json.loads(result_text)
        return json.dumps(parsed, indent=2, ensure_ascii=False)
    except (json.JSONDecodeError, ValueError):
        return json.dumps({"result": result_text}, indent=2, ensure_ascii=False)


def to_excel(results: list[dict[str, Any]], output_path: str) -> str:
    """Write a list of result dicts to a styled Excel file.

    Each dict should have keys: filename, mode, result.
    Returns the output_path on success.
    """
    rows = []
    for item in results:
        result_text = item.get("result", "")
        # Try to flatten JSON results into a readable string
        try:
            parsed = json.loads(result_text)
            if isinstance(parsed, dict):
                result_display = "\n".join(
                    f"{k}: {json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v}"
                    for k, v in parsed.items()
                )
            else:
                result_display = json.dumps(parsed, ensure_ascii=False)
        except (json.JSONDecodeError, ValueError):
            result_display = result_text

        rows.append({
            "Filename": item.get("filename", ""),
            "Mode": item.get("mode", "").capitalize(),
            "Result": result_display,
        })

    df = pd.DataFrame(rows, columns=["Filename", "Mode", "Result"])
    output_path = str(output_path)
    df.to_excel(output_path, index=False, engine="openpyxl")

    # Apply styling
    wb = load_workbook(output_path)
    ws = wb.active

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="2E4057", end_color="2E4057", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )

    # Style header row
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # Style data rows with alternating fill
    light_fill = PatternFill(start_color="F0F4F8", end_color="F0F4F8", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        fill = light_fill if row_idx % 2 == 0 else white_fill
        for cell in row:
            cell.fill = fill
            cell.border = thin_border
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    # Set column widths
    ws.column_dimensions["A"].width = 30  # Filename
    ws.column_dimensions["B"].width = 14  # Mode
    ws.column_dimensions["C"].width = 80  # Result

    # Set row heights for data rows
    for row_idx in range(2, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 80

    wb.save(output_path)
    return output_path
