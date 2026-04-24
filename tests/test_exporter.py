"""
tests/test_exporter.py — Unit tests for exporter.py
"""

import json
import pathlib
import sys
import tempfile

import pytest

sys.path.insert(0, str(pathlib.Path(__file__).parent.parent / "src"))

from exporter import to_json, to_excel


def test_to_json_valid_json_input():
    raw = '{"parties": "Alice and Bob", "date": "2026-01-01"}'
    result = to_json(raw)
    parsed = json.loads(result)
    assert parsed["parties"] == "Alice and Bob"
    assert parsed["date"] == "2026-01-01"


def test_to_json_pretty_prints():
    raw = '{"a":1,"b":2}'
    result = to_json(raw)
    # Pretty-printed JSON has newlines
    assert "\n" in result


def test_to_json_wraps_plain_text():
    raw = "This is a plain text summary."
    result = to_json(raw)
    parsed = json.loads(result)
    assert "result" in parsed
    assert parsed["result"] == raw


def test_to_json_invalid_json_wrapped():
    raw = "Not JSON at all { broken"
    result = to_json(raw)
    parsed = json.loads(result)
    assert parsed["result"] == raw


def test_to_excel_creates_file():
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        pytest.skip("openpyxl not installed")

    results = [
        {"filename": "doc1.txt", "mode": "summary", "result": "A brief summary."},
        {"filename": "doc2.txt", "mode": "extract", "result": '{"key": "value"}'},
    ]

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        tmp_path = f.name

    output = to_excel(results, tmp_path)
    assert pathlib.Path(output).exists()
    assert pathlib.Path(output).stat().st_size > 0


def test_to_excel_correct_columns():
    try:
        from openpyxl import load_workbook
    except ImportError:
        pytest.skip("openpyxl not installed")

    results = [
        {"filename": "contract.txt", "mode": "summary", "result": "Summary text here."},
    ]

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        tmp_path = f.name

    to_excel(results, tmp_path)

    wb = load_workbook(tmp_path)
    ws = wb.active
    headers = [ws.cell(row=1, column=i).value for i in range(1, 4)]
    assert headers == ["Filename", "Mode", "Result"]

    # Check data row
    assert ws.cell(row=2, column=1).value == "contract.txt"
    assert ws.cell(row=2, column=2).value == "Summary"
