"""
tests/test_parsers.py — Unit tests for parsers.py
"""

import pathlib
import sys
import tempfile

import pytest

# Allow import from src/
sys.path.insert(0, str(pathlib.Path(__file__).parent.parent / "src"))

from parsers import parse_txt, parse_file


def test_parse_txt_basic():
    with tempfile.NamedTemporaryFile(
        mode="w", suffix=".txt", delete=False, encoding="utf-8"
    ) as f:
        f.write("Hello, world!\nLine two.")
        tmp_path = f.name

    result = parse_txt(tmp_path)
    assert "Hello, world!" in result
    assert "Line two." in result


def test_parse_txt_unicode():
    with tempfile.NamedTemporaryFile(
        mode="w", suffix=".txt", delete=False, encoding="utf-8"
    ) as f:
        f.write("Ünïcödé chäracters: 中文 日本語")
        tmp_path = f.name

    result = parse_txt(tmp_path)
    assert "中文" in result


def test_parse_file_dispatches_txt():
    with tempfile.NamedTemporaryFile(
        mode="w", suffix=".txt", delete=False, encoding="utf-8"
    ) as f:
        f.write("dispatch test")
        tmp_path = f.name

    result = parse_file(tmp_path)
    assert "dispatch test" in result


def test_parse_file_unsupported_extension():
    with tempfile.NamedTemporaryFile(suffix=".xyz", delete=False) as f:
        tmp_path = f.name

    with pytest.raises(ValueError, match="Unsupported file type"):
        parse_file(tmp_path)


def test_parse_docx():
    """Create a minimal DOCX in memory and verify parsing."""
    try:
        from docx import Document
    except ImportError:
        pytest.skip("python-docx not installed")

    doc = Document()
    doc.add_paragraph("First paragraph.")
    doc.add_paragraph("Second paragraph.")

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        tmp_path = f.name

    doc.save(tmp_path)
    result = parse_file(tmp_path)
    assert "First paragraph." in result
    assert "Second paragraph." in result
